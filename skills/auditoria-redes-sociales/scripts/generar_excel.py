"""Genera el Excel de la auditoría de redes a partir del JSON de trabajo.

Uso:
    python3 generar_excel.py datos_marca.json salida.xlsx

Usa la misma estructura que calcular_metricas.py, con estos campos añadidos
(todos opcionales; lo que falte sale como "sin datos"):

Nivel raíz:
  "modo": "cliente" | "competencia",
  "insights": [...], "fortalezas": [...], "debilidades": [...],
  "urgentes": [...], "optimizacion": [...], "hipotesis": [...],
  "oportunidades": [...], "amenazas": [...], "limitaciones": [...]

Nivel red:
  "marca": "Ortiz"            (para benchmarks: varias marcas en el mismo JSON)
  "tono", "estado", "idioma", "rol", "valoracion", "cumple", "cumple_razon"

Nivel publicación:
  "tipo": "Receta / uso"      (ver references/tipos_contenido.md)
  "colab": true               (o "colab" en el formato)
  "tema", "url"

Hojas: Resumen · Publicaciones · Formatos · Tipos de contenido · Comparativa
(solo con varias marcas) · Insights · Pendientes · Metodología.
"""

import json
import os
import sys
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from calcular_metricas import DIAS, er, frecuencia_semanal, interacciones, media, mediana  # noqa: E402

TINTA = "050A0E"
PAPEL_MUDO = "F3F3F1"
DESTACADO = "E6F0E6"
LINEA = Side(style="thin", color="E2E1DD")

H_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
H_FILL = PatternFill("solid", fgColor=TINTA)
B_FONT = Font(name="Arial", size=10)
T_FONT = Font(name="Arial", bold=True, size=14)
S_FONT = Font(name="Arial", bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")


def es_colab(p):
    return bool(p.get("colab")) or "colab" in (p.get("formato") or "").lower()


def validas(posts):
    return [p for p in posts if not p.get("fijado")]


def er_frac(posts, seg):
    v = er(posts, seg)
    return None if v is None else v / 100


def r1(x):
    return None if x is None else round(x, 1)


def tabla(ws, fila, cabecera, filas, formatos=None, anchos=None):
    """Escribe una tabla con cabecera oscura. formatos: {indice_columna: number_format}."""
    for c, h in enumerate(cabecera, 1):
        cell = ws.cell(fila, c, h)
        cell.font, cell.fill, cell.alignment = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center")
    for i, valores in enumerate(filas, 1):
        for c, v in enumerate(valores, 1):
            cell = ws.cell(fila + i, c, "sin datos" if v is None else v)
            cell.font, cell.alignment, cell.border = B_FONT, WRAP, Border(bottom=LINEA)
            if formatos and (c - 1) in formatos and v is not None:
                cell.number_format = formatos[c - 1]
    if anchos:
        for c, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    return fila + len(filas) + 2


def nombre(red, datos):
    return red.get("marca") or datos.get("marca", "")


def hoja_resumen(wb, datos):
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Auditoría de redes sociales · {datos.get('marca', '')}"
    ws["A1"].font = T_FONT
    ws["A2"] = f"Fecha: {datos.get('fecha_auditoria', '')} · Modo: {datos.get('modo', 'cliente')} · Fuente: perfiles públicos (dato observado) + análisis IA"
    ws["A2"].font = B_FONT

    redes = datos["redes"]
    cab = ["Métrica"] + [f"{nombre(r, datos)} · {r['red']}" if r.get("marca") else r["red"] for r in redes]
    filas_def = [
        ("Handle", lambda r, p: r.get("handle")),
        ("URL", lambda r, p: r.get("url")),
        ("Seguidores", lambda r, p: r.get("seguidores")),
        ("Publicaciones en la muestra", lambda r, p: len(p)),
        ("Periodo de la muestra", lambda r, p: periodo(p)),
        ("Última publicación", lambda r, p: max((x["fecha"] for x in p if x.get("fecha")), default=None)),
        ("Publicaciones / semana", lambda r, p: r1(frecuencia_semanal(p))),
        ("Formato principal", lambda r, p: principal(p, "formato")),
        ("Tipo de contenido principal", lambda r, p: principal(p, "tipo")),
        ("% colaboraciones", lambda r, p: (sum(es_colab(x) for x in p) / len(p)) if p else None),
        ("Media likes", lambda r, p: r1(media(p, "likes"))),
        ("Mediana likes", lambda r, p: r1(mediana(p, "likes"))),
        ("Media comentarios", lambda r, p: r1(media(p, "comentarios"))),
        ("Media visualizaciones", lambda r, p: r1(media(p, "vistas"))),
        ("Media compartidos", lambda r, p: r1(media(p, "compartidos"))),
        ("Media guardados", lambda r, p: r1(media(p, "guardados"))),
        ("Engagement rate", lambda r, p: er_frac(p, r.get("seguidores"))),
        ("Engagement rate contenido propio", lambda r, p: er_frac([x for x in p if not es_colab(x) and "sorteo" not in (x.get("tipo") or x.get("formato") or "").lower()], r.get("seguidores"))),
        ("Idioma / mercados", lambda r, p: r.get("idioma")),
        ("Tono", lambda r, p: r.get("tono")),
        ("Estado de la cuenta", lambda r, p: r.get("estado")),
        ("Rol / enfoque", lambda r, p: r.get("rol")),
        ("Valoración", lambda r, p: r.get("valoracion")),
        ("¿Cumple la función?", lambda r, p: r.get("cumple")),
        ("Razón", lambda r, p: r.get("cumple_razon")),
    ]
    filas = []
    for etiqueta, fn in filas_def:
        filas.append([etiqueta] + [fn(r, validas(r.get("publicaciones", []))) for r in redes])
    tabla(ws, 4, cab, filas, anchos=[30] + [26] * len(redes))
    for i, (etiqueta, _) in enumerate(filas_def):
        fmt = "0.00%" if "Engagement" in etiqueta else "0%" if etiqueta.startswith("%") else "#,##0.#"
        for c in range(2, len(redes) + 2):
            ws.cell(5 + i, c).number_format = fmt
        ws.cell(5 + i, 1).font = Font(name="Arial", bold=True, size=10)
    ws.freeze_panes = "B5"


def periodo(posts):
    f = sorted(x["fecha"] for x in posts if x.get("fecha"))
    return f"{f[0]} → {f[-1]}" if f else None


def principal(posts, campo):
    c = defaultdict(int)
    for p in posts:
        if p.get(campo):
            c[p[campo]] += 1
    return max(c, key=c.get) if c else None


def hoja_publicaciones(wb, datos):
    ws = wb.create_sheet("Publicaciones")
    cab = ["Marca", "Red", "Fecha", "Día", "Formato", "Tipo de contenido", "Colaboración", "Tema",
           "Likes", "Comentarios", "Visualizaciones", "Compartidos", "Guardados", "Interacciones",
           "Engagement rate", "Fijado (excluido)", "URL"]
    filas = []
    for r in datos["redes"]:
        seg = r.get("seguidores")
        for p in r.get("publicaciones", []):
            tiene = p.get("likes") is not None or p.get("comentarios") is not None
            filas.append([
                nombre(r, datos), r["red"], p.get("fecha"),
                DIAS[date.fromisoformat(p["fecha"]).weekday()] if p.get("fecha") else None,
                p.get("formato"), p.get("tipo"), "Sí" if es_colab(p) else "No", p.get("tema"),
                p.get("likes"), p.get("comentarios"), p.get("vistas"), p.get("compartidos"), p.get("guardados"),
                interacciones(p) if tiene else None,
                interacciones(p) / seg if (tiene and seg) else None,
                "Sí" if p.get("fijado") else "No", p.get("url"),
            ])
    tabla(ws, 1, cab, filas, formatos={8: "#,##0", 9: "#,##0", 10: "#,##0", 11: "#,##0", 12: "#,##0", 13: "#,##0", 14: "0.00%"},
          anchos=[16, 12, 11, 10, 14, 22, 12, 40, 9, 11, 13, 11, 10, 12, 11, 10, 50])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(filas) + 1}"
    ws.freeze_panes = "A2"


def agregada(wb, datos, titulo, campo):
    ws = wb.create_sheet(titulo)
    cab = ["Marca", "Red", campo.capitalize() if campo == "formato" else "Tipo de contenido", "Nº", "% de la muestra",
           "Media likes", "Media comentarios", "Media visualizaciones", "Engagement rate"]
    filas = []
    for r in datos["redes"]:
        posts = validas(r.get("publicaciones", []))
        grupos = defaultdict(list)
        for p in posts:
            grupos[p.get(campo) or "sin clasificar"].append(p)
        for g, ps in sorted(grupos.items(), key=lambda kv: -len(kv[1])):
            filas.append([nombre(r, datos), r["red"], g, len(ps), len(ps) / len(posts),
                          r1(media(ps, "likes")), r1(media(ps, "comentarios")), r1(media(ps, "vistas")),
                          er_frac(ps, r.get("seguidores"))])
    tabla(ws, 1, cab, filas, formatos={4: "0%", 5: "#,##0", 6: "#,##0", 7: "#,##0", 8: "0.00%"},
          anchos=[16, 12, 24, 6, 12, 12, 14, 16, 14])
    # Destaca, por marca y red, el grupo con mayor engagement rate.
    mejores = {}
    for i, f in enumerate(filas):
        g, v = (f[0], f[1]), f[8]
        if "sorteo" in str(f[2]).lower():
            continue  # un sorteo siempre "gana" por su mecánica; no es una referencia útil
        if v is not None and (g not in mejores or v > mejores[g][1]):
            mejores[g] = (i, v)
    for i, _ in mejores.values():
        for c in range(1, len(cab) + 1):
            ws.cell(2 + i, c).fill = PatternFill("solid", fgColor=DESTACADO)
    ws.freeze_panes = "A2"
    ws.cell(len(filas) + 3, 1, "En verde: lo que mejor rinde (mayor engagement rate) en cada red, sin contar sorteos.").font = Font(name="Arial", italic=True, size=9)


def hoja_comparativa(wb, datos):
    marcas = {nombre(r, datos) for r in datos["redes"]}
    if len(marcas) < 2:
        return
    ws = wb.create_sheet("Comparativa")
    cab = ["Marca", "Red", "Seguidores", "Publicaciones / semana", "Formato principal", "Tipo principal",
           "Mediana likes", "Media comentarios", "Engagement rate", "ER contenido propio", "Valoración", "¿Cumple su referencia?"]
    filas = []
    for r in datos["redes"]:
        p = validas(r.get("publicaciones", []))
        propio = [x for x in p if not es_colab(x) and "sorteo" not in (x.get("tipo") or x.get("formato") or "").lower()]
        filas.append([nombre(r, datos), r["red"], r.get("seguidores"), r1(frecuencia_semanal(p)),
                      principal(p, "formato"), principal(p, "tipo"), r1(mediana(p, "likes")),
                      r1(media(p, "comentarios")), er_frac(p, r.get("seguidores")), er_frac(propio, r.get("seguidores")),
                      r.get("valoracion"), r.get("cumple")])
    tabla(ws, 1, cab, filas, formatos={2: "#,##0", 6: "#,##0", 7: "#,##0", 8: "0.00%", 9: "0.00%"},
          anchos=[18, 12, 12, 13, 16, 20, 12, 13, 12, 13, 12, 22])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(filas) + 1}"
    ws.freeze_panes = "C2"


def hoja_listas(wb, titulo, datos, secciones):
    ws = wb.create_sheet(titulo)
    ws.column_dimensions["A"].width = 120
    fila = 1
    for clave, etiqueta in secciones:
        items = datos.get(clave) or []
        if not items:
            continue
        ws.cell(fila, 1, etiqueta).font = S_FONT
        fila += 1
        for it in items:
            c = ws.cell(fila, 1, f"• {it}")
            c.font, c.alignment = B_FONT, WRAP
            fila += 1
        fila += 1


def hoja_metodologia(wb):
    ws = wb.create_sheet("Metodología")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    filas = [
        ("Engagement rate", "(likes + comentarios) / seguidores, calculado por publicación y promediado. Mide qué parte de la audiencia interactúa con cada post. Es la métrica de la rejilla de redes del SWIN (1.4: «engagement, rate medio y por formato»)."),
        ("ER contenido propio", "El mismo cálculo sin colaboraciones con creadores ni sorteos: mide el rendimiento del contenido que hace la propia marca."),
        ("Muestra", "Últimas 30 publicaciones o últimos 90 días (lo que sea menor). Sin sesión iniciada, Instagram solo muestra unas 12."),
        ("Exclusiones", "Las publicaciones fijadas se listan pero no entran en las medias (suelen ser antiguas y distorsionan)."),
        ("Mediana", "Valor central: describe el post 'normal' cuando un viral o un sorteo disparan la media."),
        ("Valoración", "Escala SWIN: fuerte / medio / débil / inexistente (la marca no tiene la red) / sin datos (existe pero no se pudo ver)."),
        ("¿Cumple la función?", "Estándar de auditoría SWIN: Sí / Parcialmente / No, contrastando lo medido con el rol que debería cumplir la cuenta."),
        ("Fuente", "Cifras: observadas en el perfil público en la fecha de la auditoría. Tono, estado, valoración e insights: análisis IA. Datos aportados por el cliente: se indican como tales."),
        ("Formato vs. tipo", "Formato = cómo está hecho (reel, carrusel…). Tipo de contenido = qué cuenta (receta, producto, sorteo…)."),
    ]
    tabla(ws, 1, ["Concepto", "Definición"], filas)


def generar(datos, salida):
    wb = Workbook()
    hoja_resumen(wb, datos)
    hoja_publicaciones(wb, datos)
    agregada(wb, datos, "Formatos", "formato")
    agregada(wb, datos, "Tipos de contenido", "tipo")
    hoja_comparativa(wb, datos)
    hoja_listas(wb, "Insights", datos, [
        ("insights", "Insights"), ("fortalezas", "Fortalezas (palancas)"), ("debilidades", "Debilidades críticas"),
        ("urgentes", "Urgente"), ("optimizacion", "Optimización progresiva"),
        ("oportunidades", "Oportunidades"), ("amenazas", "Amenazas"), ("hipotesis", "Hipótesis a validar con el cliente"),
    ])
    hoja_listas(wb, "Pendientes", datos, [("limitaciones", "Sin datos / limitaciones y accesos pendientes")])
    hoja_metodologia(wb)
    wb.save(salida)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as fh:
        generar(json.load(fh), sys.argv[2])
    print(f"Excel generado: {sys.argv[2]}")
