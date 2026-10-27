"""Genera el Excel del análisis de sector y competencia a partir de un JSON.

Uso:
    python3 generar_excel.py analisis_marca.json salida.xlsx

Estructura del JSON (hojas y columnas: ver references/hojas_excel.md):

{
  "titulo": "Análisis de sector y competencia · Conservas Serrats",
  "subtitulo": "Modo contexto (SWIN 1.5) · 27/10/2026",
  "hojas": [
    {
      "nombre": "Tendencias",
      "nota": "Texto opcional bajo la tabla",
      "columnas": ["Tendencia", "Qué está pasando", ...],
      "filas": [["Premiumización", "...", ...], ...]
    }
  ]
}

Los valores que empiezan por http se convierten en enlaces. Las celdas vacías o null
salen como "no encontrado". La primera hoja lleva el título y el subtítulo encima.
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TINTA = "050A0E"
LINEA = Side(style="thin", color="E2E1DD")
H_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
H_FILL = PatternFill("solid", fgColor=TINTA)
B_FONT = Font(name="Arial", size=10)
L_FONT = Font(name="Arial", size=10, color="1F4E9E", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")

# Columnas de texto largo: más anchas.
LARGAS = ("qué está pasando", "evidencia", "implicación", "propuesta", "posicionamiento", "cita", "tono",
          "presencia", "hitos", "fortalezas", "debilidades", "razón", "por qué", "descripción", "cómo",
          "qué podemos", "hallazgo", "discrepancia", "reto", "oportunidad", "justificación", "valor", "fuentes",
          "diferencial", "gama", "título")


def ancho(col, filas, i):
    nombre = col.lower()
    if any(k in nombre for k in LARGAS):
        return 45
    largo = max([len(str(col))] + [len(str(f[i])) for f in filas if i < len(f) and f[i] is not None])
    return max(10, min(28, largo + 2))


def escribir_hoja(ws, hoja, fila_ini=1):
    cols, filas = hoja["columnas"], hoja.get("filas", [])
    for c, h in enumerate(cols, 1):
        cell = ws.cell(fila_ini, c, h)
        cell.font, cell.fill = H_FONT, H_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, fila in enumerate(filas, fila_ini + 1):
        for c in range(1, len(cols) + 1):
            v = fila[c - 1] if c - 1 < len(fila) else None
            v = "no encontrado" if v in (None, "") else v
            cell = ws.cell(r, c, v)
            cell.alignment, cell.border = WRAP, Border(bottom=LINEA)
            if isinstance(v, str) and v.startswith("http") and " " not in v:
                cell.hyperlink, cell.font = v, L_FONT
            else:
                cell.font = B_FONT
    for i, col in enumerate(cols):
        ws.column_dimensions[get_column_letter(i + 1)].width = ancho(col, filas, i)
    ws.freeze_panes = ws.cell(fila_ini + 1, 2)
    if filas:
        ws.auto_filter.ref = f"A{fila_ini}:{get_column_letter(len(cols))}{fila_ini + len(filas)}"
    if hoja.get("nota"):
        n = ws.cell(fila_ini + len(filas) + 2, 1, hoja["nota"])
        n.font = Font(name="Arial", italic=True, size=9)


def generar(datos, salida):
    wb = Workbook()
    wb.remove(wb.active)
    for i, hoja in enumerate(datos["hojas"]):
        ws = wb.create_sheet(hoja["nombre"][:31])
        fila_ini = 1
        if i == 0:
            ws["A1"] = datos.get("titulo", "")
            ws["A1"].font = Font(name="Arial", bold=True, size=14)
            ws["A2"] = datos.get("subtitulo", "")
            ws["A2"].font = B_FONT
            fila_ini = 4
        escribir_hoja(ws, hoja, fila_ini)
    wb.save(salida)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as fh:
        generar(json.load(fh), sys.argv[2])
    print(f"Excel generado: {sys.argv[2]}")
